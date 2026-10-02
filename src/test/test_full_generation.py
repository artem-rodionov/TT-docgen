from datetime import date
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from docx import Document
from openpyxl import Workbook

from docgen.core import get_project_data
from docgen.data_fetcher import DataFetcher
from docgen.entities import Worker, WorkType
from docgen.orchestration import build_project_from_ui_state, generate_documents


# ======================= вспомогательное =======================

def make_docx_template(path: Path, text: str) -> None:
    """Минимальный docx с jinja-плейсхолдером docxtpl."""
    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))


def read_docx(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def make_workers_xlsx(path: Path) -> None:
    """xlsx со структурой боевого файла: 15 колонок на работника."""
    wb = Workbook()

    ws_data = wb.active
    ws_data.title = "Data"

    # Все 15 колонок должны быть заполнены — иначе openpyxl обрежет хвост
    # и get_workers_data упадёт с IndexError.
    row = [
        "Департамент производство",          # 0
        "Графический дизайнер",              # 1
        "Иванов Иван Иванович",              # 2 — ФИО (используется как ключ)
        "20.06.2002",                        # 3
        "Паспорт гражданина РФ",             # 4
        "693-788",                           # 5
        "ОВД Ленинского р-на г. Екатеринбурга",  # 6
        "15.11.2014",                        # 7
        "577471",                            # 8
        "6786",                              # 9
        "г. Москва, наб. Обводного канала, д. 133, кв. 103",  # 10
        "046123663373",                      # 11
        'АО "АЛЬФА-БАНК" №1',                # 12
        "79921665959502890126",              # 13
        "898698725",                         # 14
    ]
    ws_data.append(row)

    ws_map = wb.create_sheet("Map")
    ws_map.append(["Степашка", "Иванов Иван Иванович", None, None])

    wb.save(str(path))


def make_worker_mock(full_name: str = "Иванов Иван Иванович") -> MagicMock:
    """Worker-мок со всеми атрибутами, которые используют generators."""
    w = MagicMock(spec=Worker)
    w.full_name.return_value = full_name
    w.initials.return_value = "ИИИ"
    w.outsource = False          # → money = "1000" в generate_act
    w.is_head = False
    w.passport = MagicMock()
    w.passport.date = date(1990, 6, 20)
    return w


# ======================= фикстуры =======================

@pytest.fixture
def settings(tmp_path):
    out_dir = tmp_path / "out"
    out_dir.mkdir()

    act = tmp_path / "act.docx"
    task = tmp_path / "task.docx"
    statement = tmp_path / "statement.docx"

    make_docx_template(act, "АКТ | {{ author.full_name() }} | {{ font_typeface }} | {{ money }}")
    make_docx_template(task, "ЗАДАНИЕ | {{ font_typeface }} | инстансов: {{ font_instances_count }}")
    make_docx_template(statement, "ЗАВЕРЕНИЕ | {{ font_typeface }}")

    return {
        "api_token": "TEST_TOKEN",
        "worker_table_path": str(tmp_path / "workers.xlsx"),
        "act_template_path": str(act),
        "task_template_path": str(task),
        "statement_template_path": str(statement),
        "output_dir": str(out_dir),
    }, out_dir


@pytest.fixture
def fake_session():
    """Мок requests.Session с эмуляцией API uchet.type-tech.ru."""
    session = MagicMock()

    def fake_get(url, params=None, verify=False):
        resp = MagicMock()
        resp.raise_for_status = MagicMock()
        if url.endswith("/projects"):
            resp.json.return_value = {"projects": [{"id": 1, "name": "TestProj"}]}
        elif "/project/1/authors-summary" in url:
            resp.json.return_value = {
                "start_date": "2001-01-01",
                "end_date": "2000-01-01",
                "authors": [{
                    "name": "Степашка",
                    "position": "Контент менеджер",
                    "is_outsource": False,
                    "tasks": [
                                {
                            "name": "Описание шрифта",
                            "stage": "выход шрифта"
                            }
                        ]
                    }
                ],
            }
        else:
            raise AssertionError(f"Unexpected URL: {url}")
        return resp

    session.get.side_effect = fake_get
    return session


# ======================= тест =======================

def test_full_generation_pipeline(tmp_path, settings, fake_session, mocker):
    settings_dict, out_dir = settings

    make_workers_xlsx(Path(settings_dict["worker_table_path"]))

    p = settings_dict["worker_table_path"]
    assert Path(p).exists(), f"xlsx не создан: {p}"
    assert Path(p).suffix == ".xlsx", f"неправильное расширение: {p}"

    worker = make_worker_mock()
    mocker.patch("docgen.core.create_workers_from_map", return_value=[worker])
    mocker.patch("docgen.generators.split_into_columns", return_value=([], [], []))
    mocker.patch.object(Worker, "get_outsource_table", return_value=[])
    mocker.patch.object(Worker, "get_insource_table", return_value=[])

    def fetcher_factory(token, path):
            return DataFetcher(token, path, session=fake_session)

    project = get_project_data("TestProj", settings_dict, fetcher_factory=fetcher_factory)
    assert project.work_name == "TestProj"
    assert len(project.workers) == 1
    assert project.workers[0] is worker

    fake_font = MagicMock()
    fake_font.name_and_version.return_value = "TestFont 1.0"
    fake_font.font_instances = []

    build_project_from_ui_state(
        project,
        start_date=date(2024, 1, 15),
        end_date=date(2024, 3, 20),
        current_date=date(2024, 2, 10),
        style="Regular",
        font_path="/fake/font.glyphs",
        head_index=0,
        work_type=WorkType.Create,
        font_factory=lambda p: fake_font,
    )
    assert project.head is worker
    assert project.head.is_head is True
    assert project.font is fake_font
    assert project.style == "Regular"
    assert project.type == WorkType.Create

    progress = []
    generate_documents(
        project,
        settings_dict,
        do_task=True,
        do_statement=True,
        do_act=True,
        progress_cb=progress.append,
    )

    assert progress == [0, 33, 66, 100]

    act_files = list(out_dir.glob("act_*.docx"))
    task_files = list(out_dir.glob("task_*.docx"))
    statement_files = list(out_dir.glob("statement_*.docx"))

    assert len(act_files) == 1, f"act: {act_files}"
    assert len(task_files) == 1, f"task: {task_files}"
    assert len(statement_files) == 1, f"statement: {statement_files}"

    act_text = read_docx(act_files[0])
    assert "АКТ" in act_text
    assert "Иванов Иван Иванович" in act_text
    assert "TestFont 1.0" in act_text
    assert "1000" in act_text

    task_text = read_docx(task_files[0])
    assert "ЗАДАНИЕ" in task_text
    assert "TestFont 1.0" in task_text
    assert "инстансов: 0" in task_text

    statement_text = read_docx(statement_files[0])
    assert "ЗАВЕРЕНИЕ" in statement_text
    assert "TestFont 1.0" in statement_text


# ======================= «совсем полный» вариант =======================

@pytest.mark.integration
def test_full_generation_with_real_worker_parsing(tmp_path, settings, fake_session, mocker):
    """То же самое, но БЕЗ мока create_workers_from_map.

    Здесь уже проверяется реальный парсинг xlsx в Worker из entities.py.
    """
    settings_dict, out_dir = settings
    make_workers_xlsx(Path(settings_dict["worker_table_path"]))

    mocker.patch("docgen.generators.split_into_columns", return_value=([], [], []))

    def fetcher_factory(token, path):
        return DataFetcher(token, path, session=fake_session)

    project = get_project_data("TestProj", settings_dict, fetcher_factory=fetcher_factory)

    assert len(project.workers) == 1
    worker = project.workers[0]
    assert worker.full_name() == "Иванов Иван Иванович"
    assert worker.outsource is False

    fake_font = MagicMock()
    fake_font.name_and_version.return_value = "TestFont 1.0"
    fake_font.font_instances = []

    build_project_from_ui_state(
        project,
        start_date=date(2024, 1, 15),
        end_date=date(2024, 3, 20),
        current_date=date(2024, 2, 10),
        style="Regular",
        font_path="/fake/font.glyphs",
        head_index=0,
        work_type=WorkType.Create,
        font_factory=lambda p: fake_font,
    )

    generate_documents(
        project, settings_dict,
        do_task=True, do_statement=True, do_act=True,
    )

    assert (out_dir / "task_TestFont 1.0.docx").exists()
    assert (out_dir / "statement_TestFont 1.0.docx").exists()
    assert len(list(out_dir.glob("act_*.docx"))) == 1